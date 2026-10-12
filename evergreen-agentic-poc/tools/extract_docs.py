#!/usr/bin/env python3
"""Stand-in for CCH ProSystem fx Scan / AutoFlow: turn a client's PBC folder into machine-readable text.

    python3 tools/extract_docs.py <client_input_dir> <out_dir>

For every file under <client_input_dir> (PERM/ and PBC/):
  * PDF  -> text of every page. Pages with almost no extractable text (phone photos, handwritten notes,
            scans) are also rendered to PNG so an agent can LOOK at them with an image-capable Read tool.
  * XLSX -> every sheet as tab-separated rows.   * DOCX -> paragraphs + tables.
  * CSV / TXT / MD / EML -> copied as text.
Writes <out_dir>/<relative_path>.txt (+ page PNGs) and <out_dir>/manifest.json listing each document with page count,
characters of text and whether it needs visual review ("needs_vision": true).

Like the real scan step, this does NOT classify, de-duplicate or interpret anything - that is the intake agent's job.
"""
import json
import os
import sys

MIN_TEXT_CHARS = 40  # pages below this are treated as image-only


def pdf_to_text(path, out_base):
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(path)
    pages, images = [], []
    for i in range(len(doc)):
        page = doc[i]
        txt = page.get_textpage().get_text_range()
        pages.append(txt)
        if len(txt.strip()) < MIN_TEXT_CHARS:
            img = f"{out_base}_p{i + 1}.png"
            page.render(scale=1.4).to_pil().save(img)
            images.append(os.path.basename(img))
    body = "".join(f"\n===== page {i + 1} =====\n{t}" for i, t in enumerate(pages))
    return body, len(doc), images


def xlsx_to_text(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"\n===== sheet: {ws.title} =====")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(out), len(wb.worksheets), []


def docx_to_text(path):
    from docx import Document
    d = Document(path)
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        out.append("----- table -----")
        for row in t.rows:
            out.append("\t".join(c.text for c in row.cells))
    return "\n".join(out), 1, []


def main(src, dst):
    os.makedirs(dst, exist_ok=True)
    manifest = []
    for root, _, files in os.walk(src):
        for fn in sorted(files):
            path = os.path.join(root, fn)
            rel = os.path.relpath(path, src)
            safe_rel = rel.replace(" ", "_").replace("(", "").replace(")", "")   # shell-friendly output names
            out_base = os.path.join(dst, safe_rel)
            os.makedirs(os.path.dirname(out_base), exist_ok=True)
            ext = fn.lower().rsplit(".", 1)[-1]
            try:
                if ext == "pdf":
                    text, n, imgs = pdf_to_text(path, out_base)
                elif ext in ("xlsx", "xlsm"):
                    text, n, imgs = xlsx_to_text(path)
                elif ext == "docx":
                    text, n, imgs = docx_to_text(path)
                elif ext in ("csv", "txt", "md", "eml", "json"):
                    text, n, imgs = open(path, errors="replace").read(), 1, []
                else:
                    text, n, imgs = "", 0, []
                err = None
            except Exception as e:  # keep going; report the failure like a scan exception
                text, n, imgs, err = "", 0, [], repr(e)
            with open(out_base + ".txt", "w") as fh:
                fh.write(text)
            manifest.append({"file": rel, "text_file": safe_rel + ".txt", "type": ext, "pages": n, "text_chars": len(text.strip()),
                             "needs_vision": bool(imgs) or (ext == "pdf" and len(text.strip()) < MIN_TEXT_CHARS),
                             "page_images": imgs, "extract_error": err})
    with open(os.path.join(dst, "manifest.json"), "w") as fh:
        json.dump(manifest, fh, indent=2)
    print(f"extracted {len(manifest)} files -> {dst} "
          f"({sum(m['needs_vision'] for m in manifest)} need visual review)")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])

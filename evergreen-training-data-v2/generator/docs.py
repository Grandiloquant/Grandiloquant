"""
Document renderers for the Project Evergreen synthetic data set.

Produces realistic-looking (but clearly synthetic) PBC documents:
  * information returns (W-2, 1099 series, 1098 series, SSA-1099, K-1s, W-2G, NR-4 ...)
  * multi-page statements (consolidated 1099s, K-1 supplemental packages, bank/closing statements, agreements)
  * "scanned" / handwritten pages rendered as images inside a PDF (noise, skew, pen colour)
  * CSV / XLSX exports (bank feeds, QuickBooks P&L, rental ledgers)
  * DOCX memos (client interview memo in firm template)
  * the final return ("Accountant's Copy") PDF
"""
from __future__ import annotations

import csv
import re
import os
import random
from io import BytesIO

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, KeepTogether)

SYN = "SYNTHETIC TRAINING DATA - NOT A REAL TAX DOCUMENT"
FONT_DIR = "/usr/share/fonts/truetype"
styles = getSampleStyleSheet()
SMALL = ParagraphStyle("small", parent=styles["Normal"], fontSize=8, leading=10)
BODY = ParagraphStyle("body", parent=styles["Normal"], fontSize=9.5, leading=12.5)
H1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=15, leading=18, spaceAfter=6)
H2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11.5, leading=14, spaceBefore=8, spaceAfter=4)
H3 = ParagraphStyle("h3", parent=styles["Heading3"], fontSize=10, leading=12, spaceBefore=6, spaceAfter=2)


def ensure_dir(path):
    os.makedirs(os.path.dirname(path), exist_ok=True)


def esc(text):
    """Escape bare ampersands for reportlab Paragraph markup (keeps existing entities and tags)."""
    return re.sub(r"&(?!(?:[a-zA-Z]+|#\d+);)", "&amp;", str(text))


def money(v, dec=2):
    if v is None or v == "":
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, float) and 0 < abs(v) < 1:   # rates / ratios
        return f"{v:.4f}".rstrip("0").rstrip(".")
    if dec == 0:
        return f"{v:,.0f}" if v >= 0 else f"({-v:,.0f})"
    return f"{v:,.2f}" if v >= 0 else f"({-v:,.2f})"


# --------------------------------------------------------------------------
# Information-return style form (single page)
# --------------------------------------------------------------------------
def info_form(path, form_no, title, year="2025", payer=None, recipient=None, boxes=None, corrected=False,
              copy_label="Copy B - For Recipient", omb="", notes=None, extra_left=None, void=False,
              watermark_dup=False, account_no=None, box_cols=2):
    """payer/recipient: list of text lines. boxes: list of (box_label, description, value)."""
    ensure_dir(path)
    c = canvas.Canvas(path, pagesize=letter)
    W, H = letter
    m = 0.5 * inch
    top = H - m
    # header
    c.setFont("Helvetica-Bold", 7)
    c.setFillColor(colors.red)
    c.drawCentredString(W / 2, H - 0.3 * inch, SYN)
    c.setFillColor(colors.black)
    if corrected:
        c.setFont("Helvetica-Bold", 11)
        c.drawString(m, top - 4, "X  CORRECTED")
    if void:
        c.setFont("Helvetica-Bold", 11)
        c.drawString(m + 150, top - 4, "VOID")
    c.setFont("Helvetica-Bold", 16)
    c.drawRightString(W - m, top - 14, f"Form {form_no}")
    c.setFont("Helvetica", 8)
    c.drawRightString(W - m, top - 26, f"Tax Year {year}   {omb}")
    c.setFont("Helvetica-Bold", 11)
    c.drawString(m, top - 24, title)
    c.setFont("Helvetica", 7.5)
    c.drawString(m, top - 36, copy_label)
    y = top - 48
    left_w = 3.2 * inch
    # left blocks
    def block(lbl, lines, y, h):
        c.rect(m, y - h, left_w, h)
        c.setFont("Helvetica", 6.5)
        c.drawString(m + 3, y - 9, lbl)
        c.setFont("Courier", 8.5)
        yy = y - 20
        for ln in lines or []:
            c.drawString(m + 6, yy, str(ln)[:52])
            yy -= 10.5
        return y - h
    yl = y
    if payer:
        yl = block(payer[0], payer[1:], yl, 16 + 10.5 * max(3, len(payer) - 1) + 6)
    if recipient:
        yl = block(recipient[0], recipient[1:], yl, 16 + 10.5 * max(3, len(recipient) - 1) + 6)
    if account_no:
        yl = block("Account number (see instructions)", [account_no], yl, 30)
    for lbl, lines in (extra_left or []):
        yl = block(lbl, lines, yl, 16 + 10.5 * max(1, len(lines)) + 6)
    # right box grid
    x0 = m + left_w
    gw = W - m - x0
    cols = box_cols
    bw = gw / cols
    bh = 34
    yr = y
    boxes = boxes or []
    for i, (bl, desc, val) in enumerate(boxes):
        col = i % cols
        if col == 0 and i:
            yr -= bh
        x = x0 + col * bw
        c.rect(x, yr - bh, bw, bh)
        c.setFont("Helvetica-Bold", 6.5)
        c.drawString(x + 3, yr - 8, str(bl))
        c.setFont("Helvetica", 6.2)
        c.drawString(x + 3 + 6.5 * len(str(bl)) * 0.62 + 4, yr - 8, str(desc)[:int(bw / 3.1)])
        c.setFont("Courier-Bold", 9.5)
        sval = money(val) if isinstance(val, (int, float)) else str(val)
        c.drawRightString(x + bw - 5, yr - bh + 8, sval[:40])
    yr -= bh
    ybot = min(yl, yr) - 14
    if notes:
        c.setFont("Helvetica", 7.2)
        for ln in notes:
            for chunk in _wrap(ln, 150):
                c.drawString(m, ybot, chunk)
                ybot -= 9.5
    if watermark_dup:
        c.saveState()
        c.setFont("Helvetica-Bold", 60)
        c.setFillColor(colors.Color(.8, .8, .8, alpha=0.35))
        c.translate(W / 2, H / 2)
        c.rotate(35)
        c.drawCentredString(0, 0, watermark_dup)
        c.restoreState()
    c.setFont("Helvetica", 6)
    c.drawString(m, 0.35 * inch, f"Form {form_no} ({year})  - generated for Project Evergreen eval suite. {SYN}")
    c.showPage()
    c.save()


def _wrap(text, n):
    words, out, cur = str(text).split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 > n:
            out.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    if cur:
        out.append(cur)
    return out or [""]


# --------------------------------------------------------------------------
# Multi-section statement (consolidated 1099, K-1 package, agreements, etc.)
# --------------------------------------------------------------------------
def statement(path, title, sections, subtitle=None, header_lines=None, landscape_mode=False, footer=None):
    """sections: list of dicts {heading, para (str|list), table: [[...]], col_widths, note, pagebreak}"""
    ensure_dir(path)
    ps = landscape(letter) if landscape_mode else letter
    doc = SimpleDocTemplate(path, pagesize=ps, leftMargin=0.55 * inch, rightMargin=0.55 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=title)
    story = []
    story.append(Paragraph(f"<font color='red' size='7'>{SYN}</font>", SMALL))
    story.append(Paragraph(esc(title), H1))
    if subtitle:
        story.append(Paragraph(esc(subtitle), BODY))
    if header_lines:
        story.append(Spacer(1, 4))
        for hl in header_lines:
            story.append(Paragraph(esc(hl), BODY))
    story.append(Spacer(1, 6))
    for s in sections:
        if s.get("pagebreak"):
            story.append(PageBreak())
        if s.get("heading"):
            story.append(Paragraph(esc(s["heading"]), H2))
        if s.get("subheading"):
            story.append(Paragraph(esc(s["subheading"]), H3))
        paras = s.get("para")
        if paras:
            for p in ([paras] if isinstance(paras, str) else paras):
                story.append(Paragraph(esc(p), BODY))
                story.append(Spacer(1, 3))
        if s.get("table"):
            data = [[Paragraph(esc(cell), SMALL) if isinstance(cell, str) and len(cell) > 28 else
                     (money(cell) if isinstance(cell, (int, float)) and not isinstance(cell, bool) else cell)
                     for cell in row] for row in s["table"]]
            t = Table(data, colWidths=s.get("col_widths"), repeatRows=1 if s.get("header", True) else 0)
            st = [("FONT", (0, 0), (-1, -1), "Helvetica", 7.6),
                  ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("ALIGN", (1, 1), (-1, -1), "RIGHT")]
            if s.get("header", True):
                st += [("BACKGROUND", (0, 0), (-1, 0), colors.Color(.88, .9, .95)),
                       ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 7.4)]
            if s.get("total_row"):
                st += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 7.6), ("LINEABOVE", (0, -1), (-1, -1), 1, colors.black)]
            if s.get("left_align_cols"):
                for ci in s["left_align_cols"]:
                    st.append(("ALIGN", (ci, 0), (ci, -1), "LEFT"))
            t.setStyle(TableStyle(st))
            story.append(t)
            story.append(Spacer(1, 5))
        if s.get("note"):
            for n in ([s["note"]] if isinstance(s["note"], str) else s["note"]):
                story.append(Paragraph(f"<i>{esc(n)}</i>", SMALL))
            story.append(Spacer(1, 4))
    if footer:
        story.append(Spacer(1, 10))
        story.append(Paragraph(esc(footer), SMALL))

    def _pg(canv, d):
        canv.setFont("Helvetica", 6.5)
        canv.drawString(0.55 * inch, 0.35 * inch, f"{title} - page {d.page}   |   {SYN}")
    doc.build(story, onFirstPage=_pg, onLaterPages=_pg)


# --------------------------------------------------------------------------
# Scanned / handwritten page(s) as image PDF
# --------------------------------------------------------------------------
def _font(name, size):
    for p in [f"{FONT_DIR}/freefont/{name}", f"{FONT_DIR}/dejavu/{name}", f"{FONT_DIR}/liberation/{name}"]:
        if os.path.exists(p):
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def scanned_pages(path, pages, handwritten=True, skew=None, noise=True, seed=1, faded=False, title_stamp=None):
    """pages: list of list[str] lines. Renders to a multi-page image PDF that looks scanned."""
    ensure_dir(path)
    rnd = random.Random(seed)
    imgs = []
    for lines in pages:
        W, H = 1275, 1650  # 150 dpi letter
        bg = (250, 248, 240) if handwritten else (255, 255, 255)
        im = Image.new("RGB", (W, H), bg)
        d = ImageDraw.Draw(im)
        if handwritten:
            for yy in range(160, H - 80, 44):
                d.line([(60, yy), (W - 60, yy)], fill=(190, 205, 230), width=1)
            d.line([(130, 0), (130, H)], fill=(235, 170, 170), width=2)
            font = _font("FreeSansOblique.ttf", 27)
            ink = (25, 45, 140) if not faded else (90, 105, 170)
        else:
            font = _font("LiberationMono-Regular.ttf", 21)
            ink = (30, 30, 30) if not faded else (120, 120, 120)
        y = 120
        d.text((W - 520, 30), SYN, fill=(200, 40, 40), font=_font("DejaVuSans.ttf", 13))
        if title_stamp:
            d.text((60, 40), title_stamp, fill=(120, 0, 0), font=_font("DejaVuSans-Bold.ttf", 22))
        for ln in lines:
            x = 150 if handwritten else 80
            if handwritten:
                x += rnd.randint(-6, 10)
                jit = rnd.randint(-3, 3)
            else:
                jit = 0
            d.text((x, y + jit), ln, fill=ink, font=font)
            y += 44 if handwritten else 30
            if y > H - 90:
                break
        if noise:
            px = im.load()
            for _ in range(9000):
                xx, yy = rnd.randrange(W), rnd.randrange(H)
                g = rnd.randint(120, 200)
                px[xx, yy] = (g, g, g)
        ang = skew if skew is not None else rnd.uniform(-1.6, 1.6)
        im = im.rotate(ang, expand=False, fillcolor=(235, 235, 235), resample=Image.BICUBIC)
        if not handwritten or faded:
            im = im.filter(ImageFilter.GaussianBlur(0.6))
        imgs.append(im)
    imgs[0].save(path, "PDF", resolution=150, save_all=True, append_images=imgs[1:])


# --------------------------------------------------------------------------
# Tabular exports
# --------------------------------------------------------------------------
def write_csv(path, header, rows):
    ensure_dir(path)
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)


def write_xlsx(path, sheets):
    """sheets: dict name -> list of rows (first row header)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    ensure_dir(path)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        for i, row in enumerate(rows):
            ws.append(row)
            if i == 0:
                for cell in ws[1]:
                    cell.font = Font(bold=True)
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
    wb.save(path)


def write_docx(path, title, paragraphs, subtitle=None):
    """paragraphs: list of str or ('h', text) or ('b', [bullets]) or ('t', rows)."""
    from docx import Document
    from docx.shared import Pt
    ensure_dir(path)
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run(SYN)
    run.font.size = Pt(7)
    doc.add_heading(title, level=1)
    if subtitle:
        doc.add_paragraph(subtitle)
    for item in paragraphs:
        if isinstance(item, tuple):
            kind, val = item
            if kind == "h":
                doc.add_heading(val, level=2)
            elif kind == "b":
                for b in val:
                    doc.add_paragraph(b, style="List Bullet")
            elif kind == "t":
                t = doc.add_table(rows=0, cols=len(val[0]))
                t.style = "Table Grid"
                for row in val:
                    cells = t.add_row().cells
                    for i, v in enumerate(row):
                        cells[i].text = money(v) if isinstance(v, (int, float)) else str(v)
        else:
            doc.add_paragraph(item)
    doc.save(path)


def write_text(path, text):
    ensure_dir(path)
    with open(path, "w") as fh:
        fh.write(text)


# --------------------------------------------------------------------------
# Final return PDF (Accountant's Copy)
# --------------------------------------------------------------------------
def return_pdf(path, header, ret, state_summary=None, attachments=None, copy_label="ACCOUNTANT'S COPY"):
    ensure_dir(path)
    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.55 * inch, rightMargin=0.55 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch, title=header["title"])
    story = [Paragraph(f"<font color='red' size='7'>{SYN}</font>", SMALL),
             Paragraph(esc(header["title"]), H1),
             Paragraph(f"<b>{copy_label}</b> - Tax Year 2025 - Prepared by Evergreen Tax (synthetic firm)", BODY),
             Spacer(1, 6)]
    info = [[Paragraph(f"<b>{esc(k)}</b>", SMALL), Paragraph(esc(v), SMALL)] for k, v in header["info"]]
    t = Table(info, colWidths=[2.1 * inch, 5.1 * inch])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8.5), ("FONT", (0, 0), (0, -1), "Helvetica-Bold", 8.5),
                           ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [t, Spacer(1, 8)]
    # summary box
    v = ret.values
    summ = [["Key figures", "Amount"],
            ["Total income (1040 line 9)", money(v.get("9", 0), 0)],
            ["Adjusted gross income (line 11)", money(v.get("11", 0), 0)],
            [f"Deduction (line 12e) - {v.get('deduction_type')}", money(v.get("12e", 0), 0)],
            ["QBI deduction (line 13a)", money(v.get("13a", 0), 0)],
            ["Schedule 1-A deductions (line 13b)", money(v.get("13b", 0), 0)],
            ["Taxable income (line 15)", money(v.get("15", 0), 0)],
            ["Total tax (line 24)", money(v.get("24", 0), 0)],
            ["Total payments (line 33)", money(v.get("33", 0), 0)],
            ["REFUND (line 35a)" if v.get("refund") else "BALANCE DUE (line 37 + 38)",
             money(v.get("refund") or v.get("balance_due", 0), 0)]]
    t = Table(summ, colWidths=[4.6 * inch, 1.6 * inch])
    t.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 9), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 9),
                           ("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 9.5),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.Color(.88, .92, .88)),
                           ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("ALIGN", (1, 0), (1, -1), "RIGHT")]))
    story += [t]
    for form, lines in ret.forms.items():
        story.append(PageBreak() if form == "Form 1040" else Spacer(1, 10))
        rows = [["Line", "Description", "Amount"]]
        for ln, d, amt in lines:
            rows.append([str(ln), Paragraph(esc(d), SMALL),
                         money(amt, 0) if isinstance(amt, (int, float)) and not isinstance(amt, bool) else str(amt)])
        tt = Table(rows, colWidths=[0.8 * inch, 5.3 * inch, 1.2 * inch], repeatRows=1)
        tt.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(.9, .9, .9)),
                                ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("ALIGN", (2, 0), (2, -1), "RIGHT"),
                                ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(KeepTogether([Paragraph(esc(form), H2)]) )
        story.append(tt)
    for title, rows in (attachments or []):
        story.append(Spacer(1, 10))
        story.append(Paragraph(esc(title), H2))
        if isinstance(rows, str):
            story.append(Paragraph(esc(rows), BODY))
        else:
            data = [[Paragraph(esc(c), SMALL) if isinstance(c, str) else money(c, 0) for c in row] for row in rows]
            tt = Table(data, repeatRows=1)
            tt.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8), ("GRID", (0, 0), (-1, -1), .25, colors.grey),
                                    ("BACKGROUND", (0, 0), (-1, 0), colors.Color(.9, .9, .9)), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(tt)
    if state_summary:
        story.append(PageBreak())
        story.append(Paragraph("State / Local Returns", H1))
        for st in state_summary:
            story.append(Paragraph(esc(st["title"]), H2))
            rows = [["Line", "Description", "Amount"]] + [[a, Paragraph(esc(b), SMALL), money(c, 0) if isinstance(c, (int, float)) else str(c)] for a, b, c in st["lines"]]
            tt = Table(rows, colWidths=[0.8 * inch, 5.3 * inch, 1.2 * inch], repeatRows=1)
            tt.setStyle(TableStyle([("FONT", (0, 0), (-1, -1), "Helvetica", 8), ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
                                    ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("ALIGN", (2, 0), (2, -1), "RIGHT"),
                                    ("BACKGROUND", (0, 0), (-1, 0), colors.Color(.9, .9, .9))]))
            story.append(tt)
            if st.get("note"):
                story.append(Paragraph(f"<i>{esc(st['note'])}</i>", SMALL))

    def _pg(canv, d):
        canv.setFont("Helvetica", 6.5)
        canv.drawString(0.55 * inch, 0.35 * inch, f"{header['title']} - {copy_label} - page {d.page}  |  {SYN}")
    doc.build(story, onFirstPage=_pg, onLaterPages=_pg)
